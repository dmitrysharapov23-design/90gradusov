"""Договор «90 градусов» (ООО / ИП / физлицо): JSON заказа -> DOCX.

Формат — как в последних договорах (ЧЭЕ1610/2026, ЕДЭ1910/2026): A4, одна общая таблица характеристик,
сроки «чертежи + производство», упрощённый акт. Правила — docs/17_договоры_правила.md.
Юридический текст — contract_text.py. Номер, суммы и платежи считает код.

    .venv/bin/python app/contracts/make_contract.py data/contracts/in/<файл>.json [--out data/contracts/out]
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import contract_text as T  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data/contracts/out"

FONT, SIZE = "Century Gothic", 10
TITLE = "на изготовление и поставку металлоконструкций"
CITY = "г. Санкт-Петербург"
PAGE_W, PAGE_H, M_LEFT, M_RIGHT, M_TOP, M_BOTTOM = 21.0, 29.7, 2.41, 1.59, 2.41, 2.41  # A4, поля рыбы
TEXT_W = PAGE_W - M_LEFT - M_RIGHT

US = ("Общество с ограниченной ответственностью «90 Градусов»", ", именуемое в дальнейшем «Исполнитель», "
      "в лице Генерального директора Шарапова Дмитрия Александровича, действующего на основании Устава, с одной стороны, и ")
US_REQUISITES = [
    "ИНН/КПП: 7838125272 / 783801001",
    "Юридический адрес: 190020, Россия, г. Санкт-Петербург, наб. Обводного канала, д. 199-201, лит. Н, пом. 7",
    "Фактический адрес: 190020, Россия, г. Санкт-Петербург, наб. Обводного канала, д. 199-201, лит. Н, пом. 7",
    "р/с: 40702810855000137337",
    "Северо-Западный Банк ПАО Сбербанк",
    "к/с: 30101810500000000653",
    "БИК: 044030653",
]
US_SIGN = "Шарапов Д.А."
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]
BLANK_DATE = "«___» __________ 20___ г."

PAY_FIRST = "оплачивается Заказчиком в качестве предоплаты на основании выставленного счета"
PAY_LAST_DELIVERY = "оплачиваются Заказчиком после уведомления о готовности всех изделий, до их отгрузки и доставки на объект"
PAY_LAST_PICKUP = "оплачиваются Заказчиком после уведомления о готовности всех изделий, до их отгрузки или передачи Заказчику"
PAY_SINGLE = "оплачивается Заказчиком на основании выставленного счета"
DRAWINGS_NOTE = "После подписания Договора и Спецификации, предоставления исходных данных и поступления 70% аванса"
SERVICES_NOTE = "Дата согласовывается после уведомления о готовности изделий и с учетом готовности объекта"
# Общая таблица «2. Технические характеристики изделий»: подпись -> ключ tech.
TECH_ROWS = [("Материалы", "materials"), ("Проектные размеры", "dimensions"), ("Финиш", "finish"),
             ("Конструктив", "construction"), ("Не входит в стоимость", "exclusions")]


class ContractError(Exception):
    pass


def money(n):
    if isinstance(n, bool) or not isinstance(n, int) or n < 0:
        raise ContractError(f"сумма должна быть целым числом рублей: {n!r}")
    return f"{n:,}".replace(",", " ")


def date_words(d):
    return f"«{d:%d}» {MONTHS[d.month - 1]} {d.year} г."


def contract_number(signer, d):
    """Первые буквы фамилии, имени, отчества подписанта + ДДММ/ГГГГ. Без отчества — ошибка, не угадываем."""
    parts = [(signer.get(k) or "").strip() for k in ("last", "first", "middle")]
    if not all(parts):
        raise ContractError("для номера договора нужны фамилия, имя и отчество подписанта со стороны Заказчика "
                            "(без отчества — уточнить правило номера)")
    return "".join(p[0].upper() for p in parts) + f"{d:%d%m}/{d.year}"


def initials(s):
    return f"{s['last']} {s['first'][0]}.{s['middle'][0]}."


def dashes(x):
    """Длинные тире в данных заказа -> короткий дефис (стандарт пакета)."""
    if isinstance(x, str):
        return re.sub(r"\s*[—–]\s*", " - ", x)
    if isinstance(x, dict):
        return {k: dashes(v) for k, v in x.items()}
    if isinstance(x, list):
        return [dashes(v) for v in x]
    return x


# --- оформление ---

def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(PAGE_W), Cm(PAGE_H)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(M_LEFT), Cm(M_RIGHT), Cm(M_TOP), Cm(M_BOTTOM)
    for name in ("Normal", "List Bullet", "Table Grid"):
        st = doc.styles[name]
        st.font.name, st.font.size = FONT, Pt(SIZE)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
    pf = doc.styles["Normal"].paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.5


def run(p, text, bold=False):
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold = FONT, Pt(SIZE), bold
    return r


def para(doc, parts, align="both", spacing=1.5, style=None, keep=False):
    """parts: строка или [(текст, жирный)]."""
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = keep
    for text, bold in ([(parts, False)] if isinstance(parts, str) else parts):
        run(p, text, bold)
    return p


def heading(doc, text):
    gap(doc)
    para(doc, [(text, True)], keep=True)


def blocks(doc, rows):
    for kind, text in rows:
        if kind == "li":
            para(doc, text, style="List Bullet")
        else:
            para(doc, [(text, kind == "h")], keep=kind == "h")


def date_line(doc, d_text):
    p = para(doc, [(CITY, False), ("\t" + d_text, False)], align="left")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)


def fix_widths(table, widths):
    """Жёсткие ширины колонок: Google Docs берёт их из tblGrid и tblW, а не из ширины ячеек."""
    table.autofit = False
    tblpr = table._tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(int(Cm(sum(widths)).twips)), "w:type": "dxa"}),
                       ("w:tblLayout", {"w:type": "fixed"})):
        el = tblpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tblpr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)


def keep_rows(table):
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def cell_text(cell, lines, bold_lines=(), align="left"):
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        run(p, line, bold=i in bold_lines)


def grid(doc, header, rows, widths, align_right=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, values in enumerate([header] + rows):
        for c, v in enumerate(values):
            cell_text(t.cell(r, c), [str(v)], align="right" if (c in align_right and r) else "left")
    fix_widths(t, widths)
    keep_rows(t)
    return t


def box(doc, height_cm):
    """Пустое поле с рамкой (пункт 6 «Чертежи», замечания в акте)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles["Table Grid"]
    fix_widths(t, [TEXT_W])
    t.rows[0].height = Cm(height_cm)
    t.rows[0].height_rule = 2  # exact


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def gap(doc):
    para(doc, "")


# --- стороны ---

def customer_clause(c):
    """Заказчик в преамбуле -> [(текст, жирный)], без окончания «с другой стороны…»."""
    if c["type"] == "ООО":
        rep = c["representative"]
        acting = "действующей" if rep.get("gender") == "ж" else "действующего"
        return [(c["full_name"], True),
                (f", ОГРН {c['ogrn']}, ИНН {c['inn']}, именуемое в дальнейшем «Заказчик», в лице {rep['position_genitive']} "
                 f"{rep['name_genitive']}, {acting} на основании {rep.get('basis_genitive') or 'Устава'}", False)]
    if c["type"] == "ИП":
        return [(f"ИП {c['full_name']}", True),
                (f", ОГРНИП {c['ogrnip']}, ИНН {c['inn']}, именуемый в дальнейшем «Заказчик»", False)]
    named = "именуемая" if c.get("gender") == "ж" else "именуемый"
    return [(c["full_name"], True), (f", {named} в дальнейшем «Заказчик»", False)]


def preamble(doc, c, ending):
    para(doc, [(US[0], True), (US[1], False)] + customer_clause(c)
         + [(f", с другой стороны, совместно именуемые «Стороны», {ending}", False)])


def customer_short(c):
    if c["type"] == "ООО":
        return c.get("short_name") or c["full_name"]
    return f"ИП {c['full_name']}" if c["type"] == "ИП" else c["full_name"]


def customer_requisites(c):
    t, b = c["type"], c.get("bank") or {}
    lines = [customer_short(c)]
    if t == "ООО":
        lines += [f"ИНН/КПП: {c['inn']}" + (f" / {c['kpp']}" if c.get("kpp") else ""), f"ОГРН: {c['ogrn']}"]
    elif t == "ИП":
        lines += [f"ИНН: {c['inn']}", f"ОГРНИП: {c['ogrnip']}"]
    else:
        ps = c.get("passport") or {}
        if c.get("birth_date"):
            lines.append(f"Дата рождения: {c['birth_date']}")
        if ps.get("series") or ps.get("number"):
            lines.append(f"Паспорт РФ: серия {ps.get('series', '')} № {ps.get('number', '')}")
        if ps.get("issued_by") or ps.get("issue_date"):
            lines.append("Выдан: " + " ".join(x for x in (ps.get("issued_by"), ps.get("issue_date")) if x))
        if ps.get("department_code"):
            lines.append(f"Код подразделения: {ps['department_code']}")
    if c.get("address"):
        lines.append(("Адрес регистрации: " if t == "физлицо" else "Адрес: ") + c["address"])
    if b.get("account"):
        lines.append(f"р/с: {b['account']}")
    if b.get("name"):
        lines.append(b["name"])
    if b.get("correspondent_account"):
        lines.append(f"к/с: {b['correspondent_account']}")
    if b.get("bik"):
        lines.append(f"БИК: {b['bik']}")
    if t == "ООО":
        rep = c["representative"]
        lines.append(f"{rep['position']}: {rep['name']}")
    for k, label in (("phone", "Телефон"), ("email", "E-mail")):
        if c.get(k):
            lines.append(f"{label}: {c[k]}")
    return lines


def signatures(doc, c, signer):
    t = doc.add_table(rows=1, cols=2)
    t.style = doc.styles["Table Grid"]
    fix_widths(t, [TEXT_W / 2, TEXT_W / 2])
    cell_text(t.cell(0, 0), ["Исполнитель", "ООО «90 Градусов»", "", f"____________________/{US_SIGN}"], bold_lines=(1,))
    cell_text(t.cell(0, 1), ["Заказчик", customer_short(c), "", f"____________________/{initials(signer)}"], bold_lines=(1,))
    keep_rows(t)


def appendix_head(doc, n, number, d):
    page_break(doc)
    for line in (f"Приложение № {n} к Договору № {number}", TITLE, f"от {date_words(d)}"):
        para(doc, line, align="right", spacing=1.0)
    gap(doc)


# --- проверки и расчёты ---

def check(data):
    c = data.get("customer") or {}
    if c.get("type") not in ("ООО", "ИП", "физлицо"):
        raise ContractError("тип Заказчика: ООО, ИП или физлицо")
    need = {"ООО": ["full_name", "inn", "ogrn"], "ИП": ["full_name", "inn", "ogrnip"], "физлицо": ["full_name"]}[c["type"]]
    miss = [k for k in need if not c.get(k)]
    if c["type"] == "ООО":
        rep = c.get("representative") or {}
        miss += [f"представитель: {k}" for k in ("position", "name", "position_genitive", "name_genitive") if not rep.get(k)]
    if miss:
        raise ContractError(f"не хватает данных Заказчика: {', '.join(miss)}")
    if not data.get("items"):
        raise ContractError("нет позиций")
    for i, it in enumerate(data["items"], 1):
        if not it.get("name") or not it.get("quantity"):
            raise ContractError(f"позиция {i}: нет названия или количества")
        if it.get("unit_price") is None and it.get("total") is None:
            raise ContractError(f"позиция {i} «{it['name']}»: нет цены")
    sv = data.get("services") or {}
    if sv.get("included") and sv.get("price") is None:
        raise ContractError(f"{sv.get('label') or 'Доставка и монтаж'} входит в стоимость, но цена не указана")
    pay = data.get("payment")
    if pay and sum(x["percent"] for x in pay) != 100:
        raise ContractError(f"проценты оплаты дают {sum(x['percent'] for x in pay)}%, а не 100%")
    if not (data.get("schedule") or {}).get("production_days"):
        raise ContractError("нет срока производства")


def line_total(it):
    q = it["quantity"]
    if it.get("unit_price") is not None:
        total = it["unit_price"] * q
        if it.get("total") is not None and it["total"] != total:
            raise ContractError(f"«{it['name']}»: {q} × {it['unit_price']} = {total}, а указано {it['total']}")
        return it["unit_price"], total
    if it["total"] % q:
        raise ContractError(f"«{it['name']}»: сумма {it['total']} не делится на {q} без остатка — укажите цену за единицу")
    return it["total"] // q, it["total"]


def payment_lines(data, total):
    """3.2: этапы с суммами; остаток от деления — в последний платёж."""
    sv = data.get("services") or {}
    pay = data.get("payment") or [{"percent": 70, "text": PAY_FIRST},
                                  {"percent": 30, "text": PAY_LAST_DELIVERY if sv.get("included") else PAY_LAST_PICKUP}]
    if len(pay) == 1 and not pay[0].get("text"):
        pay[0]["text"] = PAY_SINGLE
    amounts = [total * x["percent"] // 100 for x in pay]
    amounts[-1] = total - sum(amounts[:-1])
    lines = []
    for i, (x, a) in enumerate(zip(pay, amounts)):
        last = i == len(pay) - 1
        prefix = "оставшиеся " if last and len(pay) > 1 else ""
        lines.append(f"{prefix}{x['percent']}% от общей стоимости Заказа {x['text']}, в размере {money(a)} руб.{'' if last else ';'}")
    return lines, list(zip([x["percent"] for x in pay], amounts))


# --- документ ---

def build(data, out_dir=OUT):
    data = dashes(data)
    check(data)
    d = date.fromisoformat(data.get("date") or date.today().isoformat())
    c = data["customer"]
    signer = ((c.get("representative") or {}).get("name_parts") if c["type"] == "ООО" else c.get("name_parts")) or {}
    number = contract_number(signer, d)
    dw = date_words(d)
    sv = data.get("services") or {}
    installation = bool(sv.get("included")) and sv.get("installation", True)
    sv_label = sv.get("label") or ("Доставка и монтаж" if installation else "Доставка")
    tech = data.get("tech") or {}
    warn = []

    rows, products = [], 0
    for i, it in enumerate(data["items"], 1):
        unit_price, s = line_total(it)
        products += s
        rows.append([i, it["name"], it["quantity"], it.get("unit") or "шт.", money(unit_price), money(s)])
    total = products + (sv["price"] if sv.get("included") else 0)
    if sv.get("included"):
        rows.append([len(rows) + 1, sv_label, 1, "компл.", money(sv["price"]), money(sv["price"])])
    rows.append(["", "Итого:", "", "", "", f"{money(total)} руб."])
    pay_lines, payments = payment_lines(data, total)

    doc = Document()
    setup(doc)

    # --- Договор ---
    para(doc, [(f"Договор № {number}", True)], align="center")
    para(doc, [(TITLE, True)], align="center")
    gap(doc)
    date_line(doc, dw)
    gap(doc)
    preamble(doc, c, "заключили настоящий Договор о нижеследующем.")
    warranty = data.get("warranty_months") or 12
    for kind, text in T.BODY:
        if text.startswith(T.ONLY_INDIVIDUAL) and c["type"] != "физлицо":
            continue
        if installation:
            text = T.WITH_INSTALLATION.get(text[:4], T.WITH_INSTALLATION.get(text, text))
        if kind == "h" and re.match(r"\d+\. ", text):
            heading(doc, text)
            continue
        if text.startswith("6.2. Гарантийный срок"):
            para(doc, f"6.2. Гарантийный срок составляет {warranty} месяцев с даты подписания Заказчиком акта приема-передачи.")
            continue
        blocks(doc, [(kind, text)])
        if text.startswith("3.2. Оплата производится"):
            for line in pay_lines:
                para(doc, line, style="List Bullet")
    heading(doc, "13. Реквизиты и подписи сторон")
    t = doc.add_table(rows=2, cols=2)
    t.style = doc.styles["Table Grid"]
    fix_widths(t, [TEXT_W / 2, TEXT_W / 2])
    cell_text(t.cell(0, 0), ["Исполнитель:", "ООО «90 Градусов»"] + US_REQUISITES, bold_lines=(1,))
    cell_text(t.cell(0, 1), ["Заказчик:"] + customer_requisites(c), bold_lines=(1,))
    cell_text(t.cell(1, 0), ["", "", f"____________________/{US_SIGN}"])
    cell_text(t.cell(1, 1), ["", "", f"____________________/{initials(signer)}"])
    keep_rows(t)

    # --- Приложение 1: Спецификация ---
    appendix_head(doc, 1, number, d)
    date_line(doc, f"«___» __________ {d.year} г.")
    para(doc, [("СПЕЦИФИКАЦИЯ", True)], align="center")
    gap(doc)
    preamble(doc, c, f"согласовали настоящую Спецификацию к Договору № {number} от {dw} о нижеследующем.")

    heading(doc, "1. Предмет заказа")
    grid(doc, ["№", "Наименование изделия / работ / услуг", "Кол-во", "Ед. изм.", "Цена за ед., руб.", "Сумма, руб."],
         rows, [0.8, 6.9, 1.6, 1.9, 2.6, TEXT_W - 13.8], align_right=(4, 5))

    heading(doc, "2. Технические характеристики изделий")
    listing = "; ".join(f"{(it['name'][0].lower() + it['name'][1:]) if i else it['name']} - {it['quantity']} {it.get('unit') or 'шт.'}"
                        for i, it in enumerate(data["items"]))
    if sv.get("included"):
        listing += f"; {sv_label.lower()} - 1 комплект"
    trows = [["Перечень изделий / работ", listing + "."]]
    for label, key in TECH_ROWS:
        val = tech.get(key)
        if key == "exclusions" and not val and not sv.get("included"):
            val = "Доставка; разгрузка; подъем; монтаж."
        if val:
            trows.append([label, val])
    if sv.get("included"):
        if sv.get("scope"):
            trows.append([sv_label, sv["scope"]])
        if sv.get("address"):
            trows.append([f"Адрес {'доставки и монтажа' if installation else 'доставки'}", sv["address"]])
    trows.append(["Согласование", tech.get("approval") or T.APPROVAL_DEFAULT])
    trows.append(["Стоимость", (f"Стоимость изделий - {money(products)} руб.; {sv_label.lower()} - {money(sv['price'])} руб.; "
                                f"общая стоимость Договора - {money(total)} руб." if sv.get("included")
                                else f"Общая стоимость изделий - {money(products)} руб.")])
    if not any(tech.get(k) for _, k in TECH_ROWS[:4]):
        warn.append("нет ни материалов, ни размеров, ни финиша, ни конструктива: таблица характеристик почти пустая")
    grid(doc, ["Параметр", "Содержание"], trows, [4.5, TEXT_W - 4.5])

    heading(doc, "3. Исходные данные и документация")
    blocks(doc, T.SPEC_SOURCE)

    heading(doc, "4. Сроки и этапы выполнения")
    sch = data["schedule"]
    srows = []
    if sch.get("drawings_days"):
        srows.append(["Конструкторские чертежи", f"{sch['drawings_days']} рабочих дней", DRAWINGS_NOTE])
    srows.append(["Производство", f"{sch['production_days']} рабочих дней", "После согласования всех конструкторских чертежей"])
    if sv.get("included"):
        srows.append([sv_label, "Включены" if installation else "Включена", sv.get("term") or SERVICES_NOTE])
    grid(doc, ["Этап", "Срок", "Примечание"], srows, [4.5, 3.5, TEXT_W - 8.0])
    n = 1
    if sch.get("drawings_days"):
        para(doc, f"4.{n}. Срок разработки конструкторских чертежей составляет {sch['drawings_days']} рабочих дней с даты подписания "
                  "Договора и Спецификации, предоставления Заказчиком необходимых исходных данных и поступления 70% авансового платежа.")
        n += 1
    para(doc, f"4.{n}. Срок производства составляет {sch['production_days']} рабочих дней и исчисляется с даты согласования "
              "Заказчиком всех конструкторских чертежей.")

    heading(doc, "5. Контактные данные")
    person = data.get("contact_person") or (c["representative"]["name"] if c["type"] == "ООО" else c["full_name"])
    crows = [["Контактное лицо Заказчика", person]]
    for key, label in (("contact_phone", "Телефон"), ("contact_email", "E-mail")):
        val = data.get(key) or c.get(key.replace("contact_", ""))
        if val:
            crows.append([label, val])
    if data.get("special_conditions"):
        crows.append(["Особые условия", data["special_conditions"]])
    grid(doc, ["Позиция", "Содержание"], crows, [6.0, TEXT_W - 6.0])

    heading(doc, "6. Чертежи")
    box(doc, 9.0)
    gap(doc)
    signatures(doc, c, signer)

    # --- Приложение 2: Акт ---
    appendix_head(doc, 2, number, d)
    para(doc, [("АКТ ПРИЕМА-ПЕРЕДАЧИ", True)], align="center")
    gap(doc)
    date_line(doc, BLANK_DATE)
    preamble(doc, c, f"составили настоящий Акт к Договору № {number} от {dw} о нижеследующем.")
    heading(doc, "1. Передаваемые изделия / результат работ")
    act_name = data.get("act_title") or "Комплект изделий"
    if sv.get("included") and "доставк" not in act_name.lower():
        act_name += f" с {'доставкой и монтажом' if installation else 'доставкой'}"
    grid(doc, ["№", "Наименование изделия / работ", "Кол-во", "Ед. изм.", "Характеристики / ссылка на ТЗ / проект"],
         [[1, act_name, 1, "компл.", "См. Спецификацию и согласованные конструкторские чертежи"]],
         [0.8, 6.2, 1.7, 1.9, TEXT_W - 10.6])
    gap(doc)
    blocks(doc, T.ACT_CONFIRM_NEW)
    heading(doc, "3. Замечания Заказчика")
    box(doc, 4.5)
    gap(doc)
    blocks(doc, T.ACT_WARRANTY)
    para(doc, f"уведомлен, что гарантия {warranty} месяцев действует при соблюдении правил эксплуатации, хранения, "
              "транспортировки, монтажа и допустимых нагрузок.", style="List Bullet")
    heading(doc, "5. Подписи сторон")
    signatures(doc, c, signer)

    # --- Приложение 3: Правила эксплуатации ---
    appendix_head(doc, 3, number, d)
    para(doc, [("ПРАВИЛА ЭКСПЛУАТАЦИИ ИЗДЕЛИЙ ИЗ МЕТАЛЛОКОНСТРУКЦИЙ", True)], align="center")
    blocks(doc, T.RULES)
    gap(doc)
    signatures(doc, c, signer)

    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        cell.text for tb in doc.tables for row in tb.rows for cell in row.cells)
    if re.search("[—–]", text):
        raise ContractError("в тексте осталось длинное тире")
    if re.search(r"\[[^\]]*\]|\{\{", text):
        raise ContractError("в тексте остались служебные метки [..]")

    out_dir.mkdir(parents=True, exist_ok=True)
    client = re.sub(r"[^\w\- ]+", "", customer_short(c)).strip().replace(" ", "_")[:40]
    path = out_dir / f"Договор_{number.replace('/', '-')}_{client}.docx"
    v = 2
    while path.exists():
        path = out_dir / f"Договор_{number.replace('/', '-')}_{client}_v{v}.docx"
        v += 1
    doc.save(str(path))
    return {"path": path, "number": number, "date": dw, "total": total, "payments": payments, "warnings": warn}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--out", default=str(OUT))
    a = ap.parse_args()
    data = json.loads(Path(a.input).read_text(encoding="utf-8"))
    try:
        r = build(data, Path(a.out))
    except ContractError as e:
        sys.exit(f"ОШИБКА: {e}")
    print(f"Готово: {r['path']}\nДоговор № {r['number']} от {r['date']}, итого {money(r['total'])} руб.")
    print("Оплата: " + "; ".join(f"{p}% = {money(s)} руб." for p, s in r["payments"]))
    for w in r["warnings"]:
        print("ВНИМАНИЕ:", w)


if __name__ == "__main__":
    main()
